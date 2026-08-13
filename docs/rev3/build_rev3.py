#!/usr/bin/env python3
"""Build the Robotinics third edition as a styled DOCX.

The source of truth is manuscrito_rev3.md.  The builder also creates the
technical diagrams, adds bookmarks and a static clickable table of contents,
and supports a second pass with page numbers extracted from a rendered PDF.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageChops, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
MANUSCRIPT = ROOT / "manuscrito_rev3.md"
ASSETS = ROOT / "assets"
GENERATED = ROOT / "generated"
DOC_LANGUAGE = "pt-BR"

BLUE_DARK = "123B5D"
BLUE = "176B9B"
BLUE_LIGHT = "E8F2F7"
CYAN = "2A91B8"
ORANGE = "E88624"
RED = "B42318"
RED_LIGHT = "FDECEA"
GREEN = "2E7D5B"
GREEN_LIGHT = "E9F5EF"
INK = "243746"
GRAY = "5E6C76"
GRAY_LIGHT = "F3F5F6"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9411  # A4 minus 2.2 cm margins on each side.


FIGURES_PT = {
    "diy_roadmap": {
        "caption": "Roteiro faça você mesmo: construir, testar e integrar as três partes do Robotinics.",
        "alt": "Três blocos mostram mecânica, eletrônica e software, cada um com construção e teste antes da integração.",
        "source": "Elaboração própria.",
        "width_cm": 16.1,
    },
    "architecture": {
        "caption": "Arquitetura em camadas da terceira edição.",
        "alt": "Fluxo em camadas entre interface humana, TCHATGPT, validação, Raspberry Pi, Arduino e hardware.",
        "source": "Elaboração própria.",
        "width_cm": 16.1,
    },
    "command_pipeline": {
        "caption": "Caminho supervisionado de um comando físico.",
        "alt": "Pipeline da intenção ao comando estruturado, validação, execução e telemetria.",
        "source": "Elaboração própria.",
        "width_cm": 16.1,
    },
    "safety_states": {
        "caption": "Máquina de estados operacionais do Robotinics Rev. 3.",
        "alt": "Estados DISARMED, ARMED, EXECUTING e FAULT com transições seguras.",
        "source": "Elaboração própria.",
        "width_cm": 15.5,
    },
    "robot_cad": {
        "caption": "Modelo CAD do conjunto Robotinics preservado do acervo original.",
        "alt": "Renderização CAD frontal do robô Robotinics com cabeça, tronco, braços e base móvel.",
        "source": "Acervo do projeto Robotinics.",
        "width_cm": 10.2,
    },
    "arm_cad": {
        "caption": "Conjunto mecânico do braço e da garra.",
        "alt": "Renderização CAD lateral do braço articulado e da garra do Robotinics.",
        "source": "Acervo do projeto Robotinics.",
        "width_cm": 15.8,
    },
    "base_cad": {
        "caption": "Peças superiores da base mecânica no acervo CAD do projeto.",
        "alt": "Renderização CAD das peças superiores que fecham e estruturam a base móvel.",
        "source": "Acervo do projeto Robotinics.",
        "width_cm": 15.4,
    },
    "gear_reduction": {
        "caption": "Exemplo de redução 3:1: menor velocidade e maior torque disponível no eixo de saída.",
        "alt": "Engrenagem motora de 12 dentes aciona engrenagem de saída de 36 dentes, produzindo redução de três para um.",
        "source": "Elaboração própria.",
        "width_cm": 14.8,
    },
    "power_architecture": {
        "caption": "Distribuição de energia com proteção e trilhos separados.",
        "alt": "Pack de bateria 3S protegido alimentando trilhos separados de tração, servos e lógica.",
        "source": "Elaboração própria.",
        "width_cm": 16.0,
    },
    "shield_keyed": {
        "caption": "Shield de referência para Arduino Mega com conectores carenados e polarizados.",
        "alt": "Shield preta para Arduino Mega com conectores brancos carenados, cabeçalhos azuis e identificação de portas.",
        "source": "Referência visual fornecida pelo autor; modelo comercial meramente ilustrativo.",
        "width_cm": 12.2,
    },
    "keyed_connector": {
        "caption": "Conector polarizado: trava mecânica, numeração e ordem elétrica continuam obrigatórias.",
        "alt": "Desenho de conectores carenados de três e quatro vias com chanfro, posição um e exemplos de sinais.",
        "source": "Elaboração própria.",
        "width_cm": 15.3,
    },
    "mega_wiring": {
        "caption": "Mapa de ligações derivado do firmware histórico robotinics.ino.",
        "alt": "Arduino Mega ligado a servos, ultrassons, driver de motores, sensores analógicos, Raspberry Pi, Bluetooth, GPS e I2C, com os pinos listados.",
        "source": "Elaboração própria a partir de Software/arduino/robotinics/robotinics.ino.",
        "width_cm": 16.1,
    },
    "driver_l298n": {
        "caption": "Módulo L298N empregado no protótipo histórico para o acionamento dos motores.",
        "alt": "Fotografia de módulo driver de motores L298N com dissipador, bornes e pinos de controle.",
        "source": "Acervo da edição anterior do projeto Robotinics.",
        "width_cm": 10.5,
    },
    "sensor_current": {
        "caption": "Exemplo histórico de ligação do sensor de corrente ACS712 ao Arduino.",
        "alt": "Diagrama de um módulo ACS712 ligado ao Arduino, com o caminho de corrente identificado.",
        "source": "Acervo da edição anterior do projeto Robotinics.",
        "width_cm": 9.8,
    },
    "sensor_ultrasound": {
        "caption": "Sensor ultrassônico HC-SR04 usado para medir distância.",
        "alt": "Fotografia frontal do módulo HC-SR04 com dois transdutores e quatro pinos.",
        "source": "Acervo da edição anterior do projeto Robotinics.",
        "width_cm": 8.8,
    },
    "ultrasound_principle": {
        "caption": "Princípio de medição por tempo de voo do ultrassom.",
        "alt": "Pulso ultrassônico parte do sensor, reflete em um objeto e retorna para medição do tempo de voo.",
        "source": "Acervo da edição anterior do projeto Robotinics, redesenhado editorialmente.",
        "width_cm": 13.5,
    },
    "servo_pcb": {
        "caption": "Placa controladora de servos preservada no acervo eletrônico do Robotinics.",
        "alt": "Vista da placa controladora de servo motor do projeto Robotinics.",
        "source": "Acervo do projeto Robotinics.",
        "width_cm": 10.4,
    },
    "pi5_stack": {
        "caption": "Pilha atualizada de software e percepção no Raspberry Pi 5.",
        "alt": "Câmera e serviços locais entram no Raspberry Pi 5, que executa Picamera2, OpenCV, supervisão e integração TCHATGPT antes do validador e do Arduino Mega.",
        "source": "Elaboração própria.",
        "width_cm": 16.0,
    },
    "llm_pipeline": {
        "caption": "Papel do LLM: interpretar intenção e propor uma saída estruturada, sem acionar hardware diretamente.",
        "alt": "Entrada em linguagem natural passa por um modelo de linguagem, contrato JSON, validador determinístico e somente então pode chegar ao catálogo de ações.",
        "source": "Elaboração própria.",
        "width_cm": 16.0,
    },
    "vision_pipeline": {
        "caption": "Pipeline de visão separado do controle de movimento.",
        "alt": "Câmera, captura, percepção, evento estruturado e política de segurança em sequência.",
        "source": "Elaboração própria.",
        "width_cm": 15.0,
    },
    "test_pyramid": {
        "caption": "Pirâmide de validação: muitos testes baratos sustentam poucos ensaios completos.",
        "alt": "Pirâmide com testes unitários, integração, bancada e operação no piso.",
        "source": "Elaboração própria.",
        "width_cm": 13.5,
    },
}

FIGURES_EN = {
    "diy_roadmap": {
        "caption": "Do-it-yourself roadmap: build, test, and integrate the three Robotinics parts.",
        "alt": "Three blocks show mechanics, electronics, and software, each built and tested before integration.",
        "source": "Created by the author.",
        "width_cm": 16.1,
    },
    "architecture": {
        "caption": "Layered architecture of the third edition.",
        "alt": "Layered flow among the human interface, TCHATGPT, validation, Raspberry Pi, Arduino, and hardware.",
        "source": "Created by the author.",
        "width_cm": 16.1,
    },
    "command_pipeline": {
        "caption": "Supervised path of a physical command.",
        "alt": "Pipeline from intent through a structured command, validation, execution, and telemetry.",
        "source": "Created by the author.",
        "width_cm": 16.1,
    },
    "safety_states": {
        "caption": "Robotinics Rev. 3 operating-state machine.",
        "alt": "DISARMED, ARMED, EXECUTING, and FAULT states with safe transitions.",
        "source": "Created by the author.",
        "width_cm": 15.5,
    },
    "robot_cad": {
        "caption": "CAD model of the Robotinics assembly preserved from the original archive.",
        "alt": "Front CAD rendering of the Robotinics robot with head, torso, arms, and mobile base.",
        "source": "Robotinics project archive.",
        "width_cm": 10.2,
    },
    "arm_cad": {
        "caption": "Mechanical assembly of the arm and gripper.",
        "alt": "Side CAD rendering of the articulated arm and Robotinics gripper.",
        "source": "Robotinics project archive.",
        "width_cm": 15.8,
    },
    "base_cad": {
        "caption": "Upper mechanical-base parts from the project CAD archive.",
        "alt": "CAD rendering of upper parts that close and reinforce the mobile base.",
        "source": "Robotinics project archive.",
        "width_cm": 15.4,
    },
    "gear_reduction": {
        "caption": "A 3:1 reduction example: lower speed and greater available output torque.",
        "alt": "A 12-tooth motor gear drives a 36-tooth output gear for a three-to-one reduction.",
        "source": "Created by the author.",
        "width_cm": 14.8,
    },
    "power_architecture": {
        "caption": "Power distribution with protection and separate rails.",
        "alt": "Protected 3S battery pack feeding separate traction, servo, and logic rails.",
        "source": "Created by the author.",
        "width_cm": 16.0,
    },
    "shield_keyed": {"caption": "Reference Arduino Mega shield with shrouded polarized connectors.", "alt": "Black Arduino Mega shield with white shrouded connectors, blue headers, and labeled ports.", "source": "Visual reference supplied by the author; commercial model shown for illustration only.", "width_cm": 12.2},
    "keyed_connector": {"caption": "Polarized connector: mechanical keying does not replace numbering and pin-order verification.", "alt": "Three- and four-way shrouded connectors with key, position one, and signal examples.", "source": "Created by the author.", "width_cm": 15.3},
    "mega_wiring": {"caption": "Connection map derived from the historical robotinics.ino firmware.", "alt": "Arduino Mega connected to servos, ultrasonic sensors, motor driver, analog sensors, Raspberry Pi, Bluetooth, GPS, and I2C.", "source": "Created from Software/arduino/robotinics/robotinics.ino.", "width_cm": 16.1},
    "driver_l298n": {"caption": "L298N module used in the historical prototype.", "alt": "L298N motor-driver module with heat sink, terminals, and control pins.", "source": "Robotinics previous-edition archive.", "width_cm": 10.5},
    "sensor_current": {"caption": "Historical ACS712 current-sensor connection to Arduino.", "alt": "Diagram of an ACS712 module connected to Arduino with the current path identified.", "source": "Robotinics previous-edition archive.", "width_cm": 9.8},
    "sensor_ultrasound": {"caption": "HC-SR04 ultrasonic distance sensor.", "alt": "Front view of the HC-SR04 module with two transducers and four pins.", "source": "Robotinics previous-edition archive.", "width_cm": 8.8},
    "ultrasound_principle": {"caption": "Ultrasonic time-of-flight measurement principle.", "alt": "An ultrasonic pulse reflects from an object and returns to the sensor.", "source": "Robotinics previous-edition archive, editorially redrawn.", "width_cm": 13.5},
    "servo_pcb": {"caption": "Servo controller board preserved in the Robotinics electronics archive.", "alt": "Robotinics servo controller board.", "source": "Robotinics project archive.", "width_cm": 10.4},
    "pi5_stack": {"caption": "Updated Raspberry Pi 5 software and perception stack.", "alt": "Camera and local services enter a Raspberry Pi 5 running Picamera2, OpenCV, supervision, and TCHATGPT integration before the validator and Arduino Mega.", "source": "Created by the author.", "width_cm": 16.0},
    "llm_pipeline": {"caption": "LLM role: interpret intent and propose structured output without directly driving hardware.", "alt": "Natural language passes through a language model, JSON contract, deterministic validator, and action catalog.", "source": "Created by the author.", "width_cm": 16.0},
    "vision_pipeline": {
        "caption": "Vision pipeline separated from motion control.",
        "alt": "Camera, capture, perception, structured event, and safety policy in sequence.",
        "source": "Created by the author.",
        "width_cm": 15.0,
    },
    "test_pyramid": {
        "caption": "Validation pyramid: many inexpensive tests support a few complete trials.",
        "alt": "Pyramid with unit, integration, bench, and floor-operation tests.",
        "source": "Created by the author.",
        "width_cm": 13.5,
    },
}

FIGURES = FIGURES_PT


def tr(portuguese: str, english: str) -> str:
    return english if DOC_LANGUAGE.lower().startswith("en") else portuguese


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    return re.sub(r"\s+", " ", text).strip().casefold()


def parse_source(path: Path) -> tuple[dict[str, str], list[str]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    meta: dict[str, str] = {}
    if lines and lines[0].strip() == "---":
        end = lines.index("---", 1)
        for line in lines[1:end]:
            if ":" not in line:
                continue
            key, value = line.split(":", 1)
            meta[key.strip()] = value.strip().strip('"')
        lines = lines[end + 1 :]
    return meta, lines


def heading_records(lines: list[str]) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    idx = 0
    for line_no, line in enumerate(lines):
        m = re.match(r"^(#{1,3})\s+(.+?)\s*$", line)
        if not m:
            continue
        idx += 1
        records.append(
            {
                "line": line_no,
                "level": len(m.group(1)),
                "title": m.group(2),
                "bookmark": f"sec{idx:03d}",
            }
        )
    return records


def font_path(bold: bool = False, mono: bool = False) -> str:
    candidates: list[Path]
    if mono:
        candidates = [
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationMono-Regular.ttf"),
        ]
    elif bold:
        candidates = [
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
        ]
    else:
        candidates = [
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    raise FileNotFoundError("No suitable font found")


def hexrgb(value: str) -> tuple[int, int, int]:
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def rounded_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    text: str,
    *,
    fill: str = BLUE_LIGHT,
    outline: str = BLUE,
    text_fill: str = INK,
    size: int = 38,
    radius: int = 28,
    subtitle: str | None = None,
) -> None:
    draw.rounded_rectangle(xy, radius=radius, fill=hexrgb(fill), outline=hexrgb(outline), width=5)
    x1, y1, x2, y2 = xy
    title_font = ImageFont.truetype(font_path(bold=True), size)
    subtitle_font = ImageFont.truetype(font_path(), max(24, size - 10))
    if subtitle:
        title_box = draw.textbbox((0, 0), text, font=title_font)
        sub_lines = wrap_text(draw, subtitle, subtitle_font, x2 - x1 - 50)
        sub_height = sum(draw.textbbox((0, 0), s, font=subtitle_font)[3] + 5 for s in sub_lines)
        total = (title_box[3] - title_box[1]) + 18 + sub_height
        y = y1 + (y2 - y1 - total) / 2
        draw.text(((x1 + x2) / 2, y), text, fill=hexrgb(text_fill), font=title_font, anchor="ma")
        y += title_box[3] - title_box[1] + 18
        for line in sub_lines:
            draw.text(((x1 + x2) / 2, y), line, fill=hexrgb(GRAY), font=subtitle_font, anchor="ma")
            y += draw.textbbox((0, 0), line, font=subtitle_font)[3] + 5
    else:
        lines = wrap_text(draw, text, title_font, x2 - x1 - 40)
        heights = [draw.textbbox((0, 0), s, font=title_font)[3] for s in lines]
        y = y1 + (y2 - y1 - sum(heights) - 8 * (len(lines) - 1)) / 2
        for line, height in zip(lines, heights):
            draw.text(((x1 + x2) / 2, y), line, fill=hexrgb(text_fill), font=title_font, anchor="ma")
            y += height + 8


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if current and draw.textlength(trial, font=font) > width:
            lines.append(current)
            current = word
        else:
            current = trial
    if current:
        lines.append(current)
    return lines or [""]


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = BLUE_DARK, width: int = 9) -> None:
    draw.line((start, end), fill=hexrgb(color), width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 28
    spread = 0.62
    p1 = (end[0] - length * math.cos(angle - spread), end[1] - length * math.sin(angle - spread))
    p2 = (end[0] - length * math.cos(angle + spread), end[1] - length * math.sin(angle + spread))
    draw.polygon([end, p1, p2], fill=hexrgb(color))


def canvas(width: int = 1800, height: int = 1000) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((18, 18, width - 18, height - 18), radius=30, outline=hexrgb("D4E1E8"), width=3)
    return image, draw


def save_diagram(image: Image.Image, name: str) -> Path:
    GENERATED.mkdir(parents=True, exist_ok=True)
    out = GENERATED / f"{name}.png"
    image.save(out, dpi=(220, 220), optimize=True)
    return out


def diagram_architecture() -> Path:
    im, d = canvas(1800, 1320)
    xs = (490, 1310)
    boxes = [
        ((560, 70, 1240, 205), tr("Pessoa e interface", "Person and interface"), ORANGE, tr("Texto, voz, visão e confirmação", "Text, voice, vision, and confirmation")),
        ((560, 280, 1240, 425), "TCHATGPT", BLUE_LIGHT, tr("Intenção, contexto, agentes, RAG", "Intent, context, agents, and RAG")),
        ((490, 500, 1310, 655), tr("Contrato + segurança", "Contract + safety"), GREEN_LIGHT, tr("JSON, catálogo MAN, limites e política", "JSON, MAN catalog, limits, and policy")),
        ((560, 735, 1240, 870), "Raspberry Pi", BLUE_LIGHT, tr("Gateway, serviços e supervisão", "Gateway, services, and supervision")),
        ((560, 945, 1240, 1080), "Arduino", BLUE_LIGHT, tr("Tempo real, watchdog e limites locais", "Real time, watchdog, and local limits")),
        ((490, 1150, 1310, 1270), tr("Sensores e atuadores", "Sensors and actuators"), GRAY_LIGHT, tr("Motores, servos, ultrassom e telemetria", "Motors, servos, ultrasound, and telemetry")),
    ]
    for xy, title, fill, sub in boxes:
        rounded_box(d, xy, title, fill=fill, subtitle=sub)
    for a, b in zip(boxes, boxes[1:]):
        arrow(d, ((a[0][0] + a[0][2]) // 2, a[0][3]), ((b[0][0] + b[0][2]) // 2, b[0][1]))
    return save_diagram(im, "architecture")


def diagram_diy_roadmap() -> Path:
    im, d = canvas(1800, 900)
    title_font = ImageFont.truetype(font_path(bold=True), 38)
    d.text(
        (900, 95),
        tr("Um projeto pronto, reconstruído etapa por etapa", "A finished project, rebuilt step by step"),
        fill=hexrgb(BLUE_DARK),
        font=title_font,
        anchor="ma",
    )
    cards = [
        ((55, 230, 555, 650), tr("PARTE I · MECÂNICA", "PART I · MECHANICS"), tr("dimensionar · fabricar\nmontar · testar", "size · fabricate\nassemble · test"), "FFF3E6", ORANGE),
        ((650, 230, 1150, 650), tr("PARTE II · ELETRÔNICA", "PART II · ELECTRONICS"), tr("alimentar · conectar\nmedir · testar", "power · connect\nmeasure · test"), GREEN_LIGHT, GREEN),
        ((1245, 230, 1745, 650), tr("PARTE III · SOFTWARE", "PART III · SOFTWARE"), tr("programar · integrar\nobservar · testar", "program · integrate\nobserve · test"), BLUE_LIGHT, BLUE),
    ]
    for idx, (xy, title, sub, fill, outline) in enumerate(cards, start=1):
        rounded_box(d, xy, title, fill=fill, outline=outline, subtitle=sub, size=31)
        if idx < len(cards):
            arrow(d, (xy[2], 440), (cards[idx][0][0] - 20, 440), BLUE_DARK)
    footer_font = ImageFont.truetype(font_path(bold=True), 31)
    d.text(
        (900, 780),
        tr("Integração final somente depois dos testes de cada subsistema", "Final integration only after each subsystem passes its tests"),
        fill=hexrgb(BLUE_DARK),
        font=footer_font,
        anchor="mm",
    )
    return save_diagram(im, "diy_roadmap")


def diagram_reduction() -> Path:
    im, d = canvas(1800, 950)
    small_center = (570, 480)
    large_center = (1115, 480)
    small_r = 145
    large_r = 300
    d.ellipse((small_center[0] - small_r, small_center[1] - small_r, small_center[0] + small_r, small_center[1] + small_r), fill=hexrgb(BLUE_LIGHT), outline=hexrgb(BLUE), width=12)
    d.ellipse((large_center[0] - large_r, large_center[1] - large_r, large_center[0] + large_r, large_center[1] + large_r), fill=hexrgb(GREEN_LIGHT), outline=hexrgb(GREEN), width=12)
    d.ellipse((small_center[0] - 28, small_center[1] - 28, small_center[0] + 28, small_center[1] + 28), fill=hexrgb(BLUE_DARK))
    d.ellipse((large_center[0] - 34, large_center[1] - 34, large_center[0] + 34, large_center[1] + 34), fill=hexrgb(BLUE_DARK))
    label = ImageFont.truetype(font_path(bold=True), 39)
    small = ImageFont.truetype(font_path(), 31)
    d.text(small_center, tr("Z₁ = 12\nMOTOR", "Z₁ = 12\nMOTOR"), fill=hexrgb(BLUE_DARK), font=label, anchor="mm", align="center")
    d.text(large_center, tr("Z₂ = 36\nSAÍDA", "Z₂ = 36\nOUTPUT"), fill=hexrgb(GREEN), font=label, anchor="mm", align="center")
    arrow(d, (420, 235), (530, 170), ORANGE, 10)
    arrow(d, (1375, 660), (1290, 750), GREEN, 10)
    d.text((280, 165), tr("rotação de entrada", "input rotation"), fill=hexrgb(ORANGE), font=small, anchor="mm")
    d.text((1470, 790), tr("rotação de saída", "output rotation"), fill=hexrgb(GREEN), font=small, anchor="mm")
    d.text((900, 80), tr("i = Z₂ / Z₁ = 36 / 12 = 3", "i = Z₂ / Z₁ = 36 / 12 = 3"), fill=hexrgb(BLUE_DARK), font=label, anchor="ma")
    d.text((900, 895), tr("Ideal: velocidade ÷ 3 · torque × 3  |  Real: descontar perdas e folgas", "Ideal: speed ÷ 3 · torque × 3  |  Real: allow for losses and backlash"), fill=hexrgb(BLUE_DARK), font=small, anchor="ms")
    return save_diagram(im, "gear_reduction")


def diagram_keyed_connector() -> Path:
    im, d = canvas(1800, 900)
    title = ImageFont.truetype(font_path(bold=True), 38)
    pin = ImageFont.truetype(font_path(bold=True), 30)
    body = ImageFont.truetype(font_path(), 28)
    d.text((900, 80), tr("A carcaça define a orientação; o esquema define a função", "The housing sets orientation; the schematic sets function"), fill=hexrgb(BLUE_DARK), font=title, anchor="ma")
    examples = [
        ((140, 250, 780, 650), tr("PORTA DE 3 VIAS", "3-WAY PORT"), ["GND", "V+", "SINAL"]),
        ((1020, 250, 1660, 650), tr("PORTA DE 4 VIAS", "4-WAY PORT"), ["GND", "V+", "TX/SDA", "RX/SCL"]),
    ]
    colors = [GRAY, RED, ORANGE, BLUE]
    for xy, heading, labels in examples:
        x1, y1, x2, y2 = xy
        d.rounded_rectangle(xy, radius=28, fill=hexrgb("F8FAFB"), outline=hexrgb(BLUE_DARK), width=7)
        d.polygon([(x1 + 150, y1), (x1 + 225, y1 - 55), (x2 - 225, y1 - 55), (x2 - 150, y1)], fill=hexrgb(BLUE_DARK))
        d.text(((x1 + x2) // 2, y1 + 80), heading, fill=hexrgb(BLUE_DARK), font=title, anchor="mm")
        gap = (x2 - x1 - 120) / len(labels)
        for idx, label in enumerate(labels):
            cx = int(x1 + 60 + gap * (idx + 0.5))
            cy = y1 + 245
            d.ellipse((cx - 38, cy - 38, cx + 38, cy + 38), fill=hexrgb(colors[idx]), outline=hexrgb(BLUE_DARK), width=3)
            d.text((cx, cy + 105), f"{idx + 1}", fill=hexrgb(BLUE_DARK), font=pin, anchor="mm")
            d.text((cx, cy + 165), label, fill=hexrgb(colors[idx]), font=pin, anchor="mm")
        d.text((x1 + 36, y2 - 30), tr("▲ posição 1", "▲ position 1"), fill=hexrgb(ORANGE), font=body, anchor="ls")
    d.text((900, 800), tr("Exemplo de convenção — confirme passo, família e pinagem na ficha técnica da shield adquirida", "Example convention — confirm pitch, connector family, and pinout in the purchased shield documentation"), fill=hexrgb(RED), font=body, anchor="mm")
    return save_diagram(im, "keyed_connector")


def diagram_mega_wiring() -> Path:
    im, d = canvas(1900, 1350)
    rounded_box(d, (700, 420, 1200, 930), "ARDUINO MEGA 2560", fill=BLUE_LIGHT, outline=BLUE, subtitle=tr("firmware robotinics.ino", "robotinics.ino firmware"), size=39)
    left = [
        ((55, 90, 610, 315), tr("ULTRASSOM", "ULTRASOUND"), "TRIG/ECHO: 2/3 · 40/35 · 41/42", GREEN_LIGHT),
        ((55, 390, 610, 640), tr("SENSORES ANALÓGICOS", "ANALOG SENSORS"), "A0 corrente · A1/A2/A3 acel.\nA4 gás · A6 bateria", GRAY_LIGHT),
        ((55, 745, 610, 1025), tr("DRIVER DE TRAÇÃO", "TRACTION DRIVER"), "ENA 26 · IN1 28 · IN2 30\nIN3 34 · IN4 32 · ENB 36", "FFF3E6"),
    ]
    right = [
        ((1290, 70, 1845, 370), tr("SERVOS", "SERVOS"), "6 · 8 · 9 · 10 · 11 · 44 · 46\nalimentação externa dedicada", GREEN_LIGHT),
        ((1290, 445, 1845, 730), tr("PORTAS DE COMUNICAÇÃO", "COMMUNICATION PORTS"), "Serial0/USB gateway · Serial1 18/19 BT\n16/17 RF histórico · Serial3 14/15 GPS", BLUE_LIGHT),
        ((1290, 825, 1845, 1070), tr("I²C E CONTROLADOR AUX.", "I²C AND AUX CONTROLLER"), "SDA 20 · SCL 21 · LCD 0x20\nSoftwareSerial 37/38", GRAY_LIGHT),
    ]
    for xy, title, sub, fill in left:
        rounded_box(d, xy, title, fill=fill, subtitle=sub, size=30)
        arrow(d, (xy[2], (xy[1] + xy[3]) // 2), (685, (xy[1] + xy[3]) // 2), BLUE_DARK, 7)
    for xy, title, sub, fill in right:
        rounded_box(d, xy, title, fill=fill, subtitle=sub, size=30)
        arrow(d, (1215, (xy[1] + xy[3]) // 2), (xy[0], (xy[1] + xy[3]) // 2), BLUE_DARK, 7)
    note_font = ImageFont.truetype(font_path(bold=True), 28)
    d.text((950, 1230), tr("Todos os pinos devem ser conferidos com a revisão do firmware antes da montagem do chicote", "Check every pin against the firmware revision before assembling the harness"), fill=hexrgb(RED), font=note_font, anchor="mm")
    return save_diagram(im, "mega_wiring")


def diagram_pi5_stack() -> Path:
    im, d = canvas(1800, 1060)
    rounded_box(d, (60, 120, 470, 390), tr("Câmera", "Camera"), fill=GRAY_LIGHT, subtitle=tr("Camera Module 3\nou AI Camera", "Camera Module 3\nor AI Camera"), size=34)
    rounded_box(d, (60, 660, 470, 930), tr("Áudio e rede", "Audio and network"), fill=GRAY_LIGHT, subtitle=tr("microfone, voz, Wi-Fi\nou Ethernet", "microphone, speech, Wi-Fi\nor Ethernet"), size=34)
    rounded_box(d, (625, 260, 1175, 800), "RASPBERRY PI 5", fill=BLUE_LIGHT, outline=BLUE, subtitle=tr("Raspberry Pi OS 64-bit\nPicamera2 · OpenCV\nserviço do robô · TCHATGPT", "64-bit Raspberry Pi OS\nPicamera2 · OpenCV\nrobot service · TCHATGPT"), size=39)
    rounded_box(d, (1330, 120, 1740, 390), tr("Validador", "Validator"), fill=GREEN_LIGHT, outline=GREEN, subtitle=tr("contrato, estado\ne limites", "contract, state,\nand limits"), size=34)
    rounded_box(d, (1330, 660, 1740, 930), "ARDUINO MEGA", fill="FFF3E6", outline=ORANGE, subtitle=tr("tempo real, sensores\ne atuadores", "real time, sensors,\nand actuators"), size=34)
    arrow(d, (470, 255), (625, 410), BLUE_DARK)
    arrow(d, (470, 795), (625, 650), BLUE_DARK)
    arrow(d, (1175, 410), (1330, 255), GREEN)
    arrow(d, (1535, 390), (1535, 660), ORANGE)
    note_font = ImageFont.truetype(font_path(bold=True), 28)
    d.text((900, 990), tr("Fonte recomendada 5 V / 5 A e resfriamento ativo sob carga sustentada", "Recommended 5 V / 5 A supply and active cooling under sustained load"), fill=hexrgb(BLUE_DARK), font=note_font, anchor="mm")
    return save_diagram(im, "pi5_stack")


def diagram_llm_pipeline() -> Path:
    im, d = canvas(1800, 820)
    blocks = [
        ((55, 260, 335, 535), tr("Pedido", "Request"), tr("linguagem natural", "natural language"), "FFF3E6"),
        ((410, 260, 690, 535), "LLM", tr("interpretação\nprobabilística", "probabilistic\ninterpretation"), BLUE_LIGHT),
        ((765, 260, 1045, 535), tr("Contrato", "Contract"), "JSON", GREEN_LIGHT),
        ((1120, 260, 1400, 535), tr("Validador", "Validator"), tr("MAN · estado\nlimites", "MAN · state\nlimits"), GREEN_LIGHT),
        ((1475, 260, 1755, 535), tr("Resultado", "Result"), tr("resposta ou ação\nautorizada", "reply or authorized\naction"), GRAY_LIGHT),
    ]
    for idx, (xy, title, sub, fill) in enumerate(blocks):
        rounded_box(d, xy, title, fill=fill, subtitle=sub, size=31)
        if idx < len(blocks) - 1:
            arrow(d, (xy[2], 397), (blocks[idx + 1][0][0] - 15, 397), BLUE_DARK, 7)
    f = ImageFont.truetype(font_path(bold=True), 33)
    d.text((900, 100), tr("O modelo propõe; o software determinístico autoriza", "The model proposes; deterministic software authorizes"), fill=hexrgb(BLUE_DARK), font=f, anchor="ma")
    d.text((900, 720), tr("Saída inválida, comando inexistente ou parâmetro fora da faixa → rejeitar e registrar", "Invalid output, unknown command, or out-of-range parameter → reject and log"), fill=hexrgb(RED), font=f, anchor="ms")
    return save_diagram(im, "llm_pipeline")


def diagram_command_pipeline() -> Path:
    im, d = canvas(1800, 780)
    x_positions = [55, 405, 755, 1105, 1455]
    titles = [
        (tr("Intenção", "Intent"), tr("fala ou texto", "speech or text")),
        (tr("Contrato", "Contract"), tr("JSON estrito", "strict JSON")),
        (tr("Validação", "Validation"), tr("MAN, estado, limites", "MAN, state, limits")),
        (tr("Execução", "Execution"), tr("firmware + watchdog", "firmware + watchdog")),
        (tr("Verificação", "Verification"), tr("telemetria e resultado", "telemetry and result")),
    ]
    fills = ["FFF3E6", BLUE_LIGHT, GREEN_LIGHT, BLUE_LIGHT, GRAY_LIGHT]
    for i, (title, sub) in enumerate(titles):
        x = x_positions[i]
        rounded_box(d, (x, 245, x + 290, 515), title, fill=fills[i], subtitle=sub, size=34)
        if i < len(titles) - 1:
            arrow(d, (x + 290, 380), (x_positions[i + 1] - 15, 380))
    label_font = ImageFont.truetype(font_path(bold=True), 34)
    d.text((900, 105), tr("Nenhuma etapa probabilística aciona o hardware diretamente", "No probabilistic stage directly drives hardware"), fill=hexrgb(BLUE_DARK), font=label_font, anchor="ma")
    return save_diagram(im, "command_pipeline")


def diagram_safety_states() -> Path:
    im, d = canvas(1800, 1000)
    coords = {
        "DISARMED": (90, 130, 640, 350),
        "ARMED": (1160, 130, 1710, 350),
        "EXECUTING": (1160, 650, 1710, 870),
        "FAULT": (90, 650, 640, 870),
    }
    rounded_box(d, coords["DISARMED"], "DISARMED", fill=GRAY_LIGHT, subtitle=tr("movimento bloqueado", "motion blocked"))
    rounded_box(d, coords["ARMED"], "ARMED", fill=GREEN_LIGHT, outline=GREEN, subtitle=tr("pronto dentro dos limites", "ready within limits"))
    rounded_box(d, coords["EXECUTING"], "EXECUTING", fill=BLUE_LIGHT, subtitle=tr("uma ação limitada", "one bounded action"))
    rounded_box(d, coords["FAULT"], "FAULT", fill=RED_LIGHT, outline=RED, text_fill=RED, subtitle=tr("saída segura e diagnóstico", "safe output and diagnostics"))
    arrow(d, (640, 240), (1160, 240), GREEN)
    arrow(d, (1435, 350), (1435, 650), BLUE_DARK)
    arrow(d, (1160, 760), (640, 760), RED)
    arrow(d, (365, 650), (365, 350), GRAY)
    f = ImageFont.truetype(font_path(bold=True), 29)
    d.text((900, 190), tr("autoteste + confirmação", "self-test + confirmation"), fill=hexrgb(GREEN), font=f, anchor="mm")
    d.text((1515, 500), tr("ação autorizada", "authorized action"), fill=hexrgb(BLUE_DARK), font=f, anchor="mm")
    d.text((900, 710), tr("falha, timeout ou emergência", "fault, timeout, or emergency"), fill=hexrgb(RED), font=f, anchor="mm")
    d.text((440, 500), tr("correção + reset", "correction + reset"), fill=hexrgb(GRAY), font=f, anchor="mm")
    d.text((900, 940), tr("STOP e DISARM permanecem disponíveis em qualquer estado", "STOP and DISARM remain available in every state"), fill=hexrgb(BLUE_DARK), font=f, anchor="mm")
    return save_diagram(im, "safety_states")


def diagram_power() -> Path:
    im, d = canvas(1800, 1020)
    rounded_box(d, (70, 365, 450, 650), tr("Pack 3S", "3S pack"), fill="FFF3E6", outline=ORANGE, subtitle=tr("11,1 V nominal\n12,6 V máximo", "11.1 V nominal\n12.6 V maximum"))
    rounded_box(d, (560, 350, 1050, 665), tr("Proteção e corte", "Protection and cutoff"), fill=RED_LIGHT, outline=RED, subtitle=tr("BMS + fusível + chave\n+ emergência", "BMS + fuse + switch\n+ emergency stop"))
    arrow(d, (450, 507), (560, 507), ORANGE)
    rails = [
        ((1250, 80, 1725, 310), tr("Tração", "Traction"), tr("driver e motores", "driver and motors"), BLUE_LIGHT),
        ((1250, 395, 1725, 625), tr("Servos", "Servos"), tr("regulador dedicado", "dedicated regulator"), GREEN_LIGHT),
        ((1250, 710, 1725, 940), tr("Lógica", "Logic"), tr("5 V estáveis", "stable 5 V"), GRAY_LIGHT),
    ]
    for xy, title, sub, fill in rails:
        rounded_box(d, xy, title, fill=fill, subtitle=sub)
        arrow(d, (1050, 507), (xy[0], (xy[1] + xy[3]) // 2), BLUE_DARK)
    f = ImageFont.truetype(font_path(bold=True), 28)
    d.text((900, 965), tr("Terra comum controlado; retorno de potência roteado para reduzir ruído", "Controlled common ground; power return routed to reduce noise"), fill=hexrgb(BLUE_DARK), font=f, anchor="mm")
    return save_diagram(im, "power_architecture")


def diagram_vision() -> Path:
    im, d = canvas(1800, 780)
    x_positions = [55, 405, 755, 1105, 1455]
    labels = [
        (tr("Câmera", "Camera"), tr("frame com timestamp", "timestamped frame")),
        (tr("Captura", "Capture"), tr("resolução e fila", "resolution and queue")),
        (tr("Percepção", "Perception"), tr("detecção/classificação", "detection/classification")),
        (tr("Evento", "Event"), tr("classe, confiança, posição", "class, confidence, position")),
        (tr("Política", "Policy"), tr("telemetria ou ação validada", "telemetry or validated action")),
    ]
    fills = [GRAY_LIGHT, BLUE_LIGHT, BLUE_LIGHT, GREEN_LIGHT, "FFF3E6"]
    for i, (title, sub) in enumerate(labels):
        x = x_positions[i]
        rounded_box(d, (x, 245, x + 290, 515), title, fill=fills[i], subtitle=sub, size=33)
        if i < 4:
            arrow(d, (x + 290, 380), (x_positions[i + 1] - 15, 380))
    f = ImageFont.truetype(font_path(bold=True), 34)
    d.text((900, 110), tr("Percepção informa; a política decide", "Perception informs; policy decides"), fill=hexrgb(BLUE_DARK), font=f, anchor="ma")
    return save_diagram(im, "vision_pipeline")


def diagram_tests() -> Path:
    im, d = canvas(1600, 1120)
    levels = [
        ((625, 95, 975, 275), tr("Operação no piso", "Floor test"), ORANGE),
        ((470, 320, 1130, 520), tr("Sistema em bancada", "Bench system"), BLUE),
        ((310, 565, 1290, 775), tr("Integração entre módulos", "Module integration"), CYAN),
        ((150, 820, 1450, 1035), tr("Testes unitários e simulação", "Unit tests and simulation"), GREEN),
    ]
    f = ImageFont.truetype(font_path(bold=True), 38)
    for xy, label, color in levels:
        d.polygon([(xy[0], xy[3]), (xy[0] + 90, xy[1]), (xy[2] - 90, xy[1]), (xy[2], xy[3])], fill=hexrgb(color), outline=hexrgb(WHITE))
        d.text(((xy[0] + xy[2]) / 2, (xy[1] + xy[3]) / 2), label, fill=hexrgb(WHITE), font=f, anchor="mm")
    return save_diagram(im, "test_pyramid")


def trim_image(path: Path, out: Path, *, pad: int = 40) -> Path:
    im = Image.open(path).convert("RGB")
    bg = Image.new("RGB", im.size, "white")
    diff = ImageChops.difference(im, bg).convert("L")
    diff = diff.point(lambda p: 255 if p > 16 else 0)
    bbox = diff.getbbox()
    if bbox:
        left = max(0, bbox[0] - pad)
        top = max(0, bbox[1] - pad)
        right = min(im.width, bbox[2] + pad)
        bottom = min(im.height, bbox[3] + pad)
        im = im.crop((left, top, right, bottom))
    im = ImageOps.expand(im, border=25, fill="white")
    out.parent.mkdir(parents=True, exist_ok=True)
    im.save(out, dpi=(220, 220), optimize=True)
    return out


def prepare_figures() -> dict[str, Path]:
    GENERATED.mkdir(parents=True, exist_ok=True)
    paths = {
        "diy_roadmap": diagram_diy_roadmap(),
        "architecture": diagram_architecture(),
        "command_pipeline": diagram_command_pipeline(),
        "safety_states": diagram_safety_states(),
        "gear_reduction": diagram_reduction(),
        "power_architecture": diagram_power(),
        "keyed_connector": diagram_keyed_connector(),
        "mega_wiring": diagram_mega_wiring(),
        "pi5_stack": diagram_pi5_stack(),
        "llm_pipeline": diagram_llm_pipeline(),
        "vision_pipeline": diagram_vision(),
        "test_pyramid": diagram_tests(),
        "robot_cad": trim_image(ASSETS / "robot_cad.png", GENERATED / "robot_cad.png", pad=12),
        "arm_cad": trim_image(ASSETS / "arm_cad.png", GENERATED / "arm_cad.png", pad=12),
        "base_cad": trim_image(ASSETS / "projeto" / "base_superior_cad.png", GENERATED / "base_superior_cad.png", pad=12),
        "shield_keyed": trim_image(ASSETS / "componentes" / "shield_mega_conectores_chaveados.png", GENERATED / "shield_mega_conectores_chaveados.png", pad=20),
        "driver_l298n": trim_image(ASSETS / "componentes" / "driver_l298n.png", GENERATED / "driver_l298n.png", pad=16),
        "sensor_current": trim_image(ASSETS / "componentes" / "sensor_corrente_acs712.png", GENERATED / "sensor_corrente_acs712.png", pad=16),
        "sensor_ultrasound": trim_image(ASSETS / "componentes" / "sensor_hcsr04.png", GENERATED / "sensor_hcsr04.png", pad=16),
        "ultrasound_principle": trim_image(ASSETS / "componentes" / "principio_ultrassom.png", GENERATED / "principio_ultrassom.png", pad=8),
        "servo_pcb": trim_image(ASSETS / "projeto" / "placa_controladora_servo.png", GENERATED / "placa_controladora_servo.png", pad=16),
    }
    return paths


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 70, start: int = 100, bottom: int = 70, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    # Word positions a table's visible border one start-cell margin to the left
    # of its text. Matching the indent to that margin aligns table text with
    # surrounding body text and satisfies the exact geometry audit.
    indent.set(qn("w:w"), "100")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_paragraph_border_and_shading(paragraph, color: str, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "20")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    p_pr.append(borders)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), "100")
    spacing.set(qn("w:after"), "120")


def set_language(element, language: str | None = None) -> None:
    language = language or DOC_LANGUAGE
    r_pr = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else element
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), language)
    lang.set(qn("w:eastAsia"), language)


def set_style_font(style, name: str, size: float, color: str = INK, bold: bool | None = None, italic: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    set_language(style._element.rPr)


def create_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, "Calibri", 11, INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    set_style_font(title, "Calibri", 30, BLUE_DARK, bold=True)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    subtitle = styles.add_style("Cover Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(subtitle, "Calibri", 17, BLUE, bold=False)
    subtitle.paragraph_format.space_after = Pt(14)

    eyebrow = styles.add_style("Eyebrow", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(eyebrow, "Calibri", 9.5, ORANGE, bold=True)
    eyebrow.paragraph_format.space_after = Pt(6)

    part = styles.add_style("Part Title", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(part, "Calibri", 22, BLUE_DARK, bold=True)
    part.paragraph_format.page_break_before = True
    part.paragraph_format.space_before = Pt(130)
    part.paragraph_format.space_after = Pt(18)
    part.paragraph_format.keep_with_next = True
    p_pr = part._element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    p_pr.append(outline)

    for idx, spec in {
        1: (16, BLUE_DARK, 18, 10, True),
        2: (13, BLUE, 14, 7, False),
        3: (11.5, INK, 10, 5, False),
    }.items():
        style = styles[f"Heading {idx}"]
        set_style_font(style, "Calibri", spec[0], spec[1], bold=True)
        style.paragraph_format.space_before = Pt(spec[2])
        style.paragraph_format.space_after = Pt(spec[3])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        if spec[4]:
            style.paragraph_format.page_break_before = True

    caption = styles["Caption"]
    set_style_font(caption, "Calibri", 9, GRAY, italic=True)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.keep_with_next = False

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(code, "DejaVu Sans Mono", 8.2, INK)
    code.paragraph_format.left_indent = Cm(0.25)
    code.paragraph_format.right_indent = Cm(0.15)
    code.paragraph_format.space_before = Pt(5)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.line_spacing = 1.0

    toc_heading = styles["TOC Heading"] if "TOC Heading" in styles else styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(toc_heading, "Calibri", 22, BLUE_DARK, bold=True)
    toc_heading.paragraph_format.space_after = Pt(12)

    toc1 = styles["TOC Level 1"] if "TOC Level 1" in styles else styles.add_style("TOC Level 1", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(toc1, "Calibri", 9.3, BLUE_DARK, bold=True)
    toc1.paragraph_format.space_after = Pt(1.5)
    toc1.paragraph_format.line_spacing = 1.0

    toc2 = styles["TOC Level 2"] if "TOC Level 2" in styles else styles.add_style("TOC Level 2", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(toc2, "Calibri", 8.6, INK)
    toc2.paragraph_format.left_indent = Cm(0.45)
    toc2.paragraph_format.space_after = Pt(1)
    toc2.paragraph_format.line_spacing = 1.0

    source = styles.add_style("Figure Source", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(source, "Calibri", 8, GRAY, italic=True)
    source.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_after = Pt(9)

    if "Hyperlink" not in styles:
        hyperlink = styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    else:
        hyperlink = styles["Hyperlink"]
    hyperlink.font.color.rgb = RGBColor.from_string(BLUE)
    hyperlink.font.underline = True


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.header_distance = Cm(0.9)
        section.footer_distance = Cm(0.85)


def add_field(paragraph, instruction: str, result: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = result
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def page_number_start(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def configure_header_footer(section) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("ROBOTINICS REV. 3  |  MARCELO MAURIN MARTINS")
    r.font.name = "Calibri"
    r.font.size = Pt(8)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(GRAY)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("ROBOTINICS  •  ")
    r.font.name = "Calibri"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    add_field(p, " PAGE ", "1")


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph, text: str, anchor: str, *, bold: bool = False, color: str = BLUE) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_external_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\(https?://[^)]+\)|https?://[^\s<>]+)")


def add_inline(paragraph, text: str, *, size: float | None = None, color: str | None = None) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            r = paragraph.add_run(text[pos : match.start()])
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        token = match.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        elif token.startswith("`"):
            r = paragraph.add_run(token[1:-1])
            r.font.name = "DejaVu Sans Mono"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "DejaVu Sans Mono")
            r.font.size = Pt(size or 9.2)
            r.font.color.rgb = RGBColor.from_string(BLUE_DARK)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EDF2F5")
            r._element.get_or_add_rPr().append(shd)
        elif token.startswith("["):
            link = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if link:
                add_external_hyperlink(paragraph, link.group(1), link.group(2))
        else:
            clean = token.rstrip(".,;:")
            suffix = token[len(clean) :]
            add_external_hyperlink(paragraph, clean, clean)
            if suffix:
                paragraph.add_run(suffix)
        pos = match.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def add_numbering(
    doc: Document,
    fmt: str,
    text: str,
    left: int = 540,
    hanging: int = 270,
    start_value: int = 1,
) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_value))
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    lvl.append(p_pr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)
    paragraph.paragraph_format.space_after = Pt(3 if DOC_LANGUAGE.lower().startswith("en") else 4)
    paragraph.paragraph_format.line_spacing = 1.22 if DOC_LANGUAGE.lower().startswith("en") else 1.25


def widths_for_table(rows: list[list[str]]) -> list[int]:
    cols = max(len(row) for row in rows)
    scores = [6.0] * cols
    for row in rows:
        for idx in range(cols):
            if idx >= len(row):
                continue
            cell = re.sub(r"[`*_]", "", row[idx])
            longest = max((len(word) for word in cell.split()), default=1)
            scores[idx] = max(scores[idx], min(28.0, longest * 1.3 + math.sqrt(len(cell) + 1)))
    total = sum(scores)
    widths = [max(900, int(TABLE_WIDTH_DXA * score / total)) for score in scores]
    delta = TABLE_WIDTH_DXA - sum(widths)
    widths[-1] += delta
    if widths[-1] < 800:
        need = 800 - widths[-1]
        donor = max(range(len(widths) - 1), key=lambda i: widths[i])
        widths[donor] -= need
        widths[-1] += need
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in rows[1]):
        rows = [rows[0]] + rows[2:]
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    widths = widths_for_table(rows)
    set_table_geometry(table, widths)
    for ridx, row in enumerate(rows):
        prevent_row_split(table.rows[ridx])
        if ridx == 0:
            set_repeat_table_header(table.rows[ridx])
        for cidx in range(cols):
            cell = table.cell(ridx, cidx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, BLUE_DARK if ridx == 0 else ("F7FAFC" if ridx % 2 == 0 else WHITE))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            text = row[cidx].strip() if cidx < len(row) else ""
            add_inline(p, text, size=8.1, color=WHITE if ridx == 0 else INK)
            if ridx == 0:
                for run in p.runs:
                    run.bold = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, key: str, path: Path, number: int) -> None:
    info = FIGURES[key]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Cm(float(info["width_cm"])))
    inline._inline.docPr.set("descr", str(info["alt"]))
    inline._inline.docPr.set("title", f"{tr('Figura', 'Figure')} {number}")
    caption = doc.add_paragraph(style="Caption")
    caption.add_run(f"{tr('Figura', 'Figure')} {number} - ").bold = True
    caption.add_run(str(info["caption"]))
    caption.add_run(f"\n{tr('Fonte', 'Source')}: {info['source']}")
    add_bookmark(caption, f"fig{number}", 1000 + number)


def add_callout(doc: Document, kind: str, body: str) -> None:
    labels = {"WARNING": tr("ATENÇÃO", "WARNING"), "DANGER": tr("PERIGO", "DANGER"), "NOTE": tr("NOTA", "NOTE")}
    colors = {"WARNING": ORANGE, "DANGER": RED, "NOTE": BLUE}
    fills = {"WARNING": "FFF5E8", "DANGER": RED_LIGHT, "NOTE": BLUE_LIGHT}
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.25)
    p.paragraph_format.right_indent = Cm(0.15)
    r = p.add_run(labels.get(kind, kind) + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(colors.get(kind, BLUE))
    add_inline(p, body)
    set_paragraph_border_and_shading(p, colors.get(kind, BLUE), fills.get(kind, BLUE_LIGHT))


def add_code_block(doc: Document, code: str, language: str) -> None:
    p = doc.add_paragraph(style="Code Block")
    if language:
        r = p.add_run(language.upper() + "\n")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(BLUE)
    lines = code.splitlines()
    for idx, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "DejaVu Sans Mono"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "DejaVu Sans Mono")
        if idx < len(lines) - 1:
            r.add_break()
    set_paragraph_border_and_shading(p, "CBD6DC", GRAY_LIGHT)


def add_cover(doc: Document, meta: dict[str, str], robot_path: Path) -> None:
    p = doc.add_paragraph(tr("TERCEIRA EDIÇÃO REVISTA E AMPLIADA", "THIRD REVISED AND EXPANDED EDITION"), style="Eyebrow")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(meta["title"], style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(meta["subtitle"], style="Cover Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = rule.add_run("━━━━━━━━━━━━━━━━━━━━━━━━")
    r.font.color.rgb = RGBColor.from_string(ORANGE)
    r.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    inline = p.add_run().add_picture(str(robot_path), width=Cm(13.8))
    inline._inline.docPr.set("descr", FIGURES["robot_cad"]["alt"])
    inline._inline.docPr.set("title", tr("Robotinics - modelo CAD", "Robotinics - CAD model"))
    p = doc.add_paragraph(meta["author"])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    p = doc.add_paragraph(meta["edition"] + "  •  " + meta["date"])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_imprint(doc: Document, meta: dict[str, str]) -> None:
    p = doc.add_paragraph(tr("FICHA DA EDIÇÃO", "EDITION INFORMATION"), style="Eyebrow")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(meta["title"], style="Heading 1")
    p.paragraph_format.page_break_before = False
    p = doc.add_paragraph()
    add_inline(p, f"{meta['subtitle']}\n")
    p.add_run(f"{tr('Autor e responsável pelo projeto', 'Author and project lead')}: {meta['author']}\n").bold = True
    p.add_run(f"{meta['edition']}, {meta['date']}.\n")
    p.add_run(tr("Idioma: português do Brasil.", "Language: English (United States)."))
    rows = [
        [tr("Projeto", "Project"), tr("Snapshot editorial", "Editorial snapshot")],
        ["Robotinics", "39c1b4610f1b51883a2d9fcdbccba2c672dafad3"],
        ["TCHATGPT", "15d5e1a7780088701716896cfe9fb3afe0e7b71a"],
    ]
    add_markdown_table(doc, rows)
    add_callout(
        doc,
        "WARNING",
        tr(
            "Projeto educacional e experimental. Baterias de íons de lítio, ferramentas, soldagem, motores e partes móveis exigem proteção, supervisão e componentes adequados. A IA nunca substitui as barreiras determinísticas de segurança.",
            "Educational and experimental project. Lithium-ion batteries, tools, soldering, motors, and moving parts require suitable protection, supervision, and components. AI never replaces deterministic safety barriers.",
        ),
    )
    p = doc.add_paragraph(tr("Fontes do projeto: ", "Project sources: "))
    add_external_hyperlink(p, tr("repositório Robotinics", "Robotinics repository"), "https://github.com/marcelomaurin/robotinics")
    p.add_run(tr(" e ", " and "))
    add_external_hyperlink(p, tr("repositório TCHATGPT", "TCHATGPT repository"), "https://github.com/marcelomaurin/CHATGPT")
    p.add_run(".")
    p = doc.add_paragraph(
        tr(
            "Nota editorial: o snapshot analisado do Robotinics não contém uma licença formal. A publicação pública desta edição deve ser acompanhada pela licença escolhida pelo autor.",
            "Editorial note: the examined Robotinics snapshot contains no formal license. Public release of this edition should be accompanied by the license chosen by the author.",
        )
    )


def add_dedication(doc: Document, meta: dict[str, str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(185)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text = meta.get("dedication", tr("À minha querida esposa, Daniela Machado. Amor eterno.", "To my beloved wife, Daniela Machado. Eternal love."))
    first, _, closing = text.partition(". ")
    r = p.add_run(first + ".")
    r.font.name = "Calibri"
    r.font.size = Pt(16)
    r.font.italic = True
    r.font.color.rgb = RGBColor.from_string(BLUE_DARK)
    if closing:
        p.add_run("\n\n")
        r = p.add_run(closing)
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(ORANGE)


def add_toc(doc: Document, headings: list[dict[str, object]], page_map: dict[str, int]) -> None:
    p = doc.add_paragraph(tr("Sumário", "Contents"), style="TOC Heading")
    add_bookmark(p, "TOC", 900)
    intro = doc.add_paragraph(tr("Clique em uma entrada para navegar no documento.", "Click an entry to navigate within the document."))
    intro.paragraph_format.space_after = Pt(8)
    for record in headings:
        level = int(record["level"])
        if level > 2:
            continue
        style = "TOC Level 1" if level == 1 else "TOC Level 2"
        p = doc.add_paragraph(style=style)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.35), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        add_internal_hyperlink(p, str(record["title"]), str(record["bookmark"]), bold=(level == 1), color=BLUE_DARK if level == 1 else INK)
        p.add_run("\t")
        page = page_map.get(str(record["bookmark"]))
        r = p.add_run(str(page) if page is not None else "000")
        r.font.size = Pt(8.6)
        r.font.color.rgb = RGBColor.from_string(GRAY)


def parse_table_line(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def add_content(
    doc: Document,
    lines: list[str],
    records: list[dict[str, object]],
    figure_paths: dict[str, Path],
) -> None:
    heading_by_line = {int(item["line"]): item for item in records}
    i = 0
    figure_number = 0
    bullet_num_id = add_numbering(doc, "bullet", "•")
    active_decimal_num: int | None = None
    previous_was_number = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            active_decimal_num = None
            previous_was_number = False
            i += 1
            continue

        record = heading_by_line.get(i)
        if record:
            level = int(record["level"])
            title = str(record["title"])
            if level == 1 and (title.startswith("Parte ") or title.startswith("Part ")):
                p = doc.add_paragraph(title, style="Part Title")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                deco = doc.add_paragraph("ROBOTINICS REV. 3")
                deco.paragraph_format.space_after = Pt(0)
                for run in deco.runs:
                    run.font.size = Pt(10)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(ORANGE)
            else:
                p = doc.add_paragraph(title, style=f"Heading {level}")
            add_bookmark(p, str(record["bookmark"]), 100 + int(str(record["bookmark"])[3:]))
            i += 1
            continue

        figure_match = re.fullmatch(r"\[FIGURE:([a-z0-9_]+)\]", stripped)
        if figure_match:
            key = figure_match.group(1)
            figure_number += 1
            add_figure(doc, key, figure_paths[key], figure_number)
            i += 1
            continue

        if stripped.startswith("> [!"):
            m = re.match(r"> \[!([A-Z]+)\]", stripped)
            kind = m.group(1) if m else "NOTE"
            body: list[str] = []
            i += 1
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                body.append(lines[i].lstrip()[1:].strip())
                i += 1
            add_callout(doc, kind, " ".join(body))
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            add_code_block(doc, "\n".join(code_lines), language)
            continue

        if stripped.startswith("|"):
            table_lines: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(parse_table_line(lines[i]))
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        bullet = re.match(r"^[-*]\s+(.+)", stripped)
        if bullet:
            p = doc.add_paragraph()
            apply_numbering(p, bullet_num_id)
            add_inline(p, bullet.group(1))
            i += 1
            previous_was_number = False
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)", stripped)
        if numbered:
            if not previous_was_number or active_decimal_num is None:
                active_decimal_num = add_numbering(doc, "decimal", "%1.", start_value=int(numbered.group(1)))
            p = doc.add_paragraph()
            apply_numbering(p, active_decimal_num)
            add_inline(p, numbered.group(2))
            i += 1
            previous_was_number = True
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1
        active_decimal_num = None
        previous_was_number = False


def build_docx(output: Path, page_map_path: Path | None = None) -> None:
    meta, lines = parse_source(MANUSCRIPT)
    records = heading_records(lines)
    page_map: dict[str, int] = {}
    if page_map_path and page_map_path.exists():
        page_map = {str(k): int(v) for k, v in json.loads(page_map_path.read_text(encoding="utf-8")).items()}
    figures = prepare_figures()

    doc = Document()
    configure_sections(doc)
    create_styles(doc)
    doc.core_properties.title = meta["title"]
    doc.core_properties.subject = meta["subtitle"]
    doc.core_properties.author = meta["author"]
    doc.core_properties.keywords = tr(
        "Robotinics, robótica, Arduino, Raspberry Pi, TCHATGPT, inteligência artificial",
        "Robotinics, robotics, Arduino, Raspberry Pi, TCHATGPT, artificial intelligence",
    )
    doc.core_properties.comments = tr("Terceira edição revista e ampliada.", "Third revised and expanded edition.")
    doc.core_properties.created = datetime(2026, 8, 12, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 8, 12, tzinfo=timezone.utc)

    cover_section = doc.sections[0]
    configure_sections(doc)
    add_cover(doc, meta, figures["robot_cad"])

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_sections(doc)
    page_number_start(body_section, 1)
    configure_header_footer(body_section)
    add_imprint(doc, meta)
    doc.add_page_break()
    add_dedication(doc, meta)
    doc.add_page_break()
    add_toc(doc, records, page_map)
    add_content(doc, lines, records, figures)

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    compat = settings.find(qn("w:compat"))
    if compat is not None:
        balance = OxmlElement("w:doNotExpandShiftReturn")
        compat.append(balance)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def extract_page_map(pdf_path: Path, output: Path) -> None:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required for page mapping") from exc
    _meta, lines = parse_source(MANUSCRIPT)
    records = heading_records(lines)
    reader = PdfReader(str(pdf_path))
    pages = [normalize(page.extract_text() or "") for page in reader.pages]
    marker = normalize(
        tr(
            "O Robotinics nasceu como um projeto multidisciplinar",
            "Robotinics began as a multidisciplinary project",
        )
    )
    body_start = next((idx for idx, text in enumerate(pages) if marker in text), None)
    if body_start is None:
        raise RuntimeError("Could not locate the beginning of the manuscript in the PDF")
    mapping: dict[str, int] = {}
    missing: list[str] = []
    for record in records:
        if int(record["level"]) > 2:
            continue
        needle = normalize(str(record["title"]))
        found = next((idx for idx in range(body_start, len(pages)) if needle in pages[idx]), None)
        if found is None:
            missing.append(str(record["title"]))
            continue
        mapping[str(record["bookmark"])] = found  # Cover is physical page 1; body numbering starts at 1 on physical page 2.
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(mapping, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Mapped {len(mapping)} headings; body begins on physical page {body_start + 1}.")
    if missing:
        print("Missing headings:")
        for title in missing:
            print(f"- {title}")


def main() -> None:
    global MANUSCRIPT, DOC_LANGUAGE, FIGURES
    parser = argparse.ArgumentParser()
    parser.add_argument("--manuscript", type=Path, default=MANUSCRIPT)
    parser.add_argument("--language", default=None)
    sub = parser.add_subparsers(dest="command", required=True)
    build = sub.add_parser("build")
    build.add_argument("--output", type=Path, required=True)
    build.add_argument("--page-map", type=Path)
    mapping = sub.add_parser("map")
    mapping.add_argument("--pdf", type=Path, required=True)
    mapping.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    MANUSCRIPT = args.manuscript.resolve()
    if args.language:
        DOC_LANGUAGE = args.language
    else:
        meta, _lines = parse_source(MANUSCRIPT)
        DOC_LANGUAGE = meta.get("language", "pt-BR")
    FIGURES = FIGURES_EN if DOC_LANGUAGE.lower().startswith("en") else FIGURES_PT
    if args.command == "build":
        build_docx(args.output, args.page_map)
    else:
        extract_page_map(args.pdf, args.output)


if __name__ == "__main__":
    main()
